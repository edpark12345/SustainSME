import openpyxl
p = r"C:\Users\claudia.daniel\Downloads\SustainSME_handover\backend\_export_test.xlsx"
wb = openpyxl.load_workbook(p)
print("SHEETS:", wb.sheetnames)
ws = wb["Calculations"]
print("HEADER:", [c.value for c in ws[1]])
print("DATA ROWS:", ws.max_row - 1)
for r in range(2, min(ws.max_row, 8) + 1):
    print("  ", ws.cell(row=r, column=1).value, "|", ws.cell(row=r, column=3).value,
          "|", ws.cell(row=r, column=7).value, "|", ws.cell(row=r, column=8).value, ws.cell(row=r, column=9).value)
print("--- SUMMARY ---")
sm = wb["Summary"]
for r in range(3, 12):
    a = sm.cell(row=r, column=1).value
    b = sm.cell(row=r, column=2).value
    if a:
        print(" ", a, "=", b)
print("--- FACTOR LIBRARY ---")
fl = wb["Factor library"]
for r in range(1, fl.max_row + 1):
    print(" ", [fl.cell(row=r, column=c).value for c in (1, 4, 5)])
