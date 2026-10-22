#!/usr/bin/env python3
"""
export_xlsx.py — build a transparent, screening-grade Excel workbook of the Scope 1, 2 & 3 calculation.

Screening-grade / transparent — NOT an assured or audit-ready inventory (a disclaimer to that
effect is written onto the Summary sheet). One workbook a bank can review / share:
  - Summary: totals that tie back, boundary, factor-source register, warnings.
  - Calculations: one row per activity line with the explicit formula
    (quantity x factor = tCO2e), tier, factor value/unit/source, full working.
  - Calculation steps: every sub-calculation broken out row by row (input →
    factor → each intermediate step → calculation → result) for full transparency.
  - Factor library: the datasets used, with source + year (traceability).

Built from the deterministic CalculationResponse, so the numbers match the app exactly.
"""
from __future__ import annotations

import io
import re
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from schema import ActivityBundle, CalculationResponse

_HDR_FILL = PatternFill("solid", fgColor="A100FF")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="A100FF")
_SUBTITLE_FONT = Font(bold=True, size=11, color="A100FF")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_STEP_FILL = PatternFill("solid", fgColor="F3E8FF")


def _split_qty(raw: str):
    m = re.match(r"^([\d.,]+)\s*(.*)$", raw or "")
    if not m:
        return None, (raw or "")
    return float(m.group(1).replace(",", "")), m.group(2)


def _fmt_qty(qty: float) -> str:
    """Thousands-separated, but keep up to 3 decimals so non-integer inputs (e.g.
    1.2 kg refill, 0.840 kg leaked) reconcile with the shown result."""
    return f"{qty:,.3f}".rstrip("0").rstrip(".")


def _formula(r) -> str:
    qty, unit = _split_qty(r.raw_input)
    if r.factor_value is None or qty is None:
        return f"{r.raw_input} -> {r.co2e_t:,.2f} tCO2e"
    div = " / 1000" if (r.factor_unit or "").lower().startswith("kg") else ""
    return f"{_fmt_qty(qty)} {unit} x {r.factor_value:g} {r.factor_unit}{div} = {r.co2e_t:,.2f} tCO2e"


def _header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT


def build_xlsx(bundle: ActivityBundle, calc: CalculationResponse, meta: dict) -> bytes:
    wb = Workbook()

    # ---------- Summary ----------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SustainSME - Scope 1, 2 & 3 Emissions"
    ws["A1"].font = _TITLE_FONT

    ws["A2"] = (
        "SCREENING-GRADE / TRANSPARENT ESTIMATE — for internal review and directional insight only. "
        "NOT an audited, verified or assurance-ready inventory. Do not use as audit proof or for "
        "compliance sign-off without independent verification of the emission factors, reporting "
        "boundary and underlying data quality."
    )
    ws["A2"].font = Font(italic=True, color="B00000")
    ws["A2"].alignment = _WRAP
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 48

    yr = bundle.boundary.reporting_year
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    sites = ", ".join(f"{s.name} ({s.country})" for s in bundle.boundary.sites) or "-"
    summary = [
        ("Generated", now),
        ("Reporting year", yr if yr is not None else "-"),
        ("Sites / boundary", sites),
        ("", ""),
        ("Scope 1 total (tCO2e)", round(calc.scope1_total_t, 3)),
        ("Scope 2 - location-based (tCO2e)", round(calc.scope2_location_t, 3)),
        ("Scope 2 - market-based (tCO2e)", round(calc.scope2_market_t, 3)),
        ("Scope 3 - spend-based (tCO2e)", round(calc.scope3_total_t, 3)),
        ("Total Scope 1 + 2 (location) (tCO2e)", round(calc.scope1_total_t + calc.scope2_location_t, 3)),
        ("Total Scope 1 + 2 + 3 (location) (tCO2e)",
         round(calc.scope1_total_t + calc.scope2_location_t + calc.scope3_total_t, 3)),
    ]
    r = 3
    for label, val in summary:
        c1 = ws.cell(row=r, column=1, value=label)
        if label:
            c1.font = _BOLD
        ws.cell(row=r, column=2, value=val)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Emission-factor sources").font = _BOLD
    r += 1
    for key in ("fuel", "grid", "residual_mix", "gwp", "leak_rates", "eui", "scope3_spend", "waste", "scope3_activity"):
        m = meta.get(key, {})
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=f"{m.get('source', '?')} ({m.get('year', '?')})")
        r += 1

    if calc.warnings:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = _BOLD
        r += 1
        for w in calc.warnings:
            ws.cell(row=r, column=1, value=w)
            r += 1

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 72

    # ---------- Calculations (one row per activity line) ----------
    ws2 = wb.create_sheet("Calculations")
    headers = ["Scope", "Category", "Activity", "Input", "Factor value", "Factor unit",
               "Calculation", "tCO2e", "Tier", "Factor source", "Working / assumptions"]
    ws2.append(headers)
    _header(ws2, 1, len(headers))
    row_i = 2
    for res in calc.results:
        working = "\n".join(res.assumptions) if res.assumptions else ""
        ws2.append([
            res.scope, res.category, res.activity_label, res.raw_input,
            res.factor_value, res.factor_unit, _formula(res), round(res.co2e_t, 4),
            res.tier, res.factor_source, working,
        ])
        # Wrap the free-text columns and size the row to the working text.
        for col in (3, 7, 10, 11):
            ws2.cell(row=row_i, column=col).alignment = _WRAP
        lines = max(1, len(res.assumptions), (len(_formula(res)) // 46) + 1)
        ws2.row_dimensions[row_i].height = 15 * min(lines, 8)
        row_i += 1
    ws2.freeze_panes = "A2"
    for i, w in enumerate([7, 26, 30, 24, 12, 18, 46, 12, 11, 40, 60], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---------- Calculation steps (every sub-calculation, row by row) ----------
    ws4 = wb.create_sheet("Calculation steps")
    ws4["A1"] = "Every calculation, step by step"
    ws4["A1"].font = _TITLE_FONT
    hdr = ["#", "Scope", "Activity", "Step", "Working"]
    ws4.append([])  # spacer row 2
    ws4.append(hdr)
    _header(ws4, 3, len(hdr))
    sr = 4
    for n, res in enumerate(calc.results, start=1):
        # Block header for the activity.
        ws4.cell(row=sr, column=1, value=n)
        ws4.cell(row=sr, column=2, value=res.scope)
        c = ws4.cell(row=sr, column=3, value=f"{res.category} — {res.activity_label}")
        c.font = _SUBTITLE_FONT
        for col in range(1, 6):
            ws4.cell(row=sr, column=col).fill = _STEP_FILL
        sr += 1

        def step(label: str, detail):
            nonlocal sr
            ws4.cell(row=sr, column=4, value=label).font = _BOLD
            cell = ws4.cell(row=sr, column=5, value=detail)
            cell.alignment = _WRAP
            sr += 1

        step("Input", res.raw_input)
        if res.factor_value is not None:
            step("Factor", f"{res.factor_value:g} {res.factor_unit or ''}  (source: {res.factor_source or '-'})")
        for j, a in enumerate(res.assumptions, start=1):
            step(f"Step {j}", a)
        step("Calculation", _formula(res))
        step("Result", f"{round(res.co2e_t, 4):,} tCO2e  (Scope {res.scope}, tier: {res.tier})")
        sr += 1  # blank separator

    for i, w in enumerate([5, 7, 46, 14, 80], start=1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = "A4"

    # ---------- Factor library ----------
    ws3 = wb.create_sheet("Factor library")
    ws3.append(["Dataset", "Metric", "Unit", "Source", "Year"])
    _header(ws3, 1, 5)
    for key in ("fuel", "grid", "residual_mix", "gwp", "leak_rates", "eui", "scope3_spend", "waste", "scope3_activity"):
        m = meta.get(key, {})
        ws3.append([key, m.get("metric", ""), m.get("unit", ""), m.get("source", ""), m.get("year", "")])
    for i, w in enumerate([14, 46, 18, 62, 8], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
