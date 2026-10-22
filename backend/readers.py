#!/usr/bin/env python3
"""
readers.py — deterministic extraction of text from client documents.

No LLM, no token cost: turn xlsx / csv / text-PDF / docx into compact text/markdown
that the LLM step then structures. Scanned/image PDFs fall back to OCR if the
optional deps (pytesseract + pdf2image) are installed; otherwise a clear marker
is returned so the caller can surface "needs vision/OCR".
"""
from __future__ import annotations

from pathlib import Path

MAX_XLSX_ROWS = 600       # cap rows per sheet (asset lists can be huge but sparse)
MAX_PDF_CHARS = 200_000


def read_any(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _read_xlsx(path)
    if ext == ".csv":
        return _truncate(path.read_text(encoding="utf-8", errors="replace"))
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext in (".txt", ".md"):
        return _truncate(path.read_text(encoding="utf-8", errors="replace"))
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
        return f"[IMAGE FILE {path.name} — needs vision/OCR; not text-extractable]"
    return f"[UNSUPPORTED FILE {path.name} ({ext})]"


def _truncate(s: str, limit: int = MAX_PDF_CHARS) -> str:
    return s if len(s) <= limit else s[:limit] + "\n[... truncated ...]"


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows_out: list[str] = []
        n = 0
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if c is None else str(c) for c in row]
            if not any(cells):
                continue
            rows_out.append(" | ".join(cells).rstrip(" |"))
            n += 1
            if n >= MAX_XLSX_ROWS:
                rows_out.append(f"[... sheet truncated at {MAX_XLSX_ROWS} non-empty rows ...]")
                break
        if rows_out:
            parts.append(f"## Sheet: {ws.title}\n" + "\n".join(rows_out))
    wb.close()
    return "\n\n".join(parts) if parts else "[empty workbook]"


def _read_pdf(path: Path) -> str:
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        try:
            from pypdf import PdfReader
            text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
        except Exception as e:
            return f"[PDF read error for {path.name}: {e}]"
    if text.strip():
        return _truncate(text)
    # No selectable text → likely a scan. Try OCR if available.
    ocr = _try_ocr(path)
    return ocr or f"[SCANNED/IMAGE PDF {path.name} — no selectable text; needs vision/OCR]"


def _try_ocr(path: Path) -> str | None:
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except Exception:
        return None
    try:
        pages = convert_from_path(str(path), dpi=200)
        return _truncate("\n".join(pytesseract.image_to_string(img) for img in pages))
    except Exception:
        return None


def _read_docx(path: Path) -> str:
    """Deterministic .docx -> Markdown (headings, lists, tables)."""
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception as e:
        return f"[docx read error for {path.name}: {e}]"

    doc = docx.Document(str(path))
    out: list[str] = []

    def emit_para(p):
        txt = " ".join(p.text.split())
        if not txt:
            return
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(int(digits), 6) if digits else 2
            out.append("#" * level + " " + txt)
        elif "list" in style:
            out.append("- " + txt)
        else:
            out.append(txt)

    def emit_table(t):
        for row in t.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if any(cells):
                out.append("| " + " | ".join(cells) + " |")

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            emit_para(Paragraph(child, doc))
        elif tag == "tbl":
            emit_table(Table(child, doc))
    return _truncate("\n".join(out)) if out else "[empty document]"
